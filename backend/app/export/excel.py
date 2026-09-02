"""Export ranked top candidates to an Excel workbook."""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["Rank", "Name", "Email", "Phone", "Experience (yrs)", "Score",
           "Status", "Matched Skills", "Missing Skills", "Notes"]
WIDTHS = [6, 24, 30, 20, 16, 8, 14, 40, 40, 40]


def build_excel(candidates: list[dict], job_title: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Shortlist"

    # Title row
    ws.cell(row=1, column=1,
            value=f"TalentLens shortlist — {job_title or 'Untitled role'}").font = Font(
        bold=True, size=13)
    ws.cell(row=2, column=1,
            value=f"Generated {datetime.now():%Y-%m-%d %H:%M} · scores assist human review, "
                  f"they are not hiring decisions").font = Font(italic=True, color="666666")

    header_row = 4
    header_fill = PatternFill("solid", fgColor="111111")
    header_font = Font(bold=True, color="F5C518")
    for col, head in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=head)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, c in enumerate(candidates, start=1):
        r = header_row + i
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=c.get("name", ""))
        ws.cell(row=r, column=3, value=c.get("email", ""))
        ws.cell(row=r, column=4, value=c.get("phone", ""))
        ws.cell(row=r, column=5, value=c.get("experience_years", 0))
        ws.cell(row=r, column=6, value=c.get("score", 0))
        ws.cell(row=r, column=7, value=(c.get("status") or "new").title())
        ws.cell(row=r, column=8, value=", ".join(c.get("matched_skills", [])))
        ws.cell(row=r, column=9, value=", ".join(c.get("missing_skills", [])))
        ws.cell(row=r, column=10, value=c.get("notes") or "")

    for col, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
